import csv
from datetime import date
from io import BytesIO, StringIO
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import APIRouter, Depends, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models import Reading, Tariff
from app.services.calculator import TariffSet, calculate_month

router = APIRouter(prefix="/api/export", tags=["export"])

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Имена такие же, как в интерфейсе приложения.
RESOURCE_NAMES = {
    "xvs": "ХВС (холодная вода)",
    "gvs": "ГВС (горячая вода)",
    "electric": "Электричество",
    "rent": "Аренда",
}

MONTHS_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]

# ---------------------------------------------------------------- данные

async def _collect(db: AsyncSession) -> dict[str, tuple[list[str], list[list]]]:
    """Собрать все таблицы выгрузки: (заголовки, строки) по разделам."""
    readings = list(
        (await db.execute(select(Reading).order_by(Reading.period.asc()))).scalars().all()
    )
    tariffs = list(
        (
            await db.execute(
                select(Tariff).order_by(Tariff.effective_from.asc(), Tariff.resource_type.asc())
            )
        )
        .scalars()
        .all()
    )

    # --- Показания счётчиков ---
    reading_headers = ["Период", "Год", "Месяц", "ХВС, м³", "ГВС, м³", "Электричество, кВт·ч"]
    reading_rows = [
        [
            r.period,
            r.period.year,
            _month_name(r.period),
            round(r.xvs_value, 3),
            round(r.gvs_value, 3),
            round(r.electric_value, 2),
        ]
        for r in readings
    ]

    # --- Тарифы ---
    tariff_headers = ["Ресурс", "Ставка, ₽", "Действует с", "Действует по"]
    tariff_rows = [
        [
            RESOURCE_NAMES.get(t.resource_type, t.resource_type),
            round(t.rate, 4),
            t.effective_from,
            t.effective_to or "",
        ]
        for t in tariffs
    ]

    # --- Расходы по месяцам (пересчёт тот же, что и в сводке) ---
    month_headers = [
        "Период", "Год", "Месяц",
        "Расход ХВС, м³", "Расход ГВС, м³", "Расход эл-ва, кВт·ч",
        "ХВС, ₽", "ГВС, ₽", "Эл-во, ₽", "Аренда, ₽", "Итого, ₽",
    ]
    month_rows: list[list] = []

    # Первое показание служит базовой точкой (как в сводке за год),
    # поэтому расчёт начинается со второго показания в истории.
    for i in range(1, len(readings)):
        prev_r = readings[i - 1]
        curr_r = readings[i]

        tariff = TariffSet(
            xvs_rate=_rate_for(tariffs, "xvs", curr_r.period),
            gvs_rate=_rate_for(tariffs, "gvs", curr_r.period),
            electric_rate=_rate_for(tariffs, "electric", curr_r.period),
            rent=_rate_for(tariffs, "rent", curr_r.period),
        )
        calc = calculate_month(
            current_xvs=curr_r.xvs_value,
            current_gvs=curr_r.gvs_value,
            current_electric=curr_r.electric_value,
            previous_xvs=prev_r.xvs_value,
            previous_gvs=prev_r.gvs_value,
            previous_electric=prev_r.electric_value,
            tariff=tariff,
        )
        month_rows.append([
            curr_r.period,
            curr_r.period.year,
            _month_name(curr_r.period),
            round(calc.xvs_consumption, 3),
            round(calc.gvs_consumption, 3),
            round(calc.electric_consumption, 2),
            calc.xvs_cost,
            calc.gvs_cost,
            calc.electric_cost,
            round(calc.rent, 2),
            calc.total,
        ])

    # --- Итоги по годам ---
    total_headers = ["Год", "Месяцев", "ХВС, ₽", "ГВС, ₽", "Эл-во, ₽", "Аренда, ₽", "Итого, ₽"]
    year_sums: dict[int, list[float]] = {}
    for row in month_rows:
        y = row[1]
        acc = year_sums.setdefault(y, [0.0, 0.0, 0.0, 0.0, 0.0, 0.0])
        acc[0] += 1
        acc[1] += row[6]
        acc[2] += row[7]
        acc[3] += row[8]
        acc[4] += row[9]
        acc[5] += row[10]

    total_rows: list[list] = []
    grand = [0.0] * 6
    for y in sorted(year_sums):
        s = [round(v, 2) for v in year_sums[y]]
        total_rows.append([y, int(s[0]), s[1], s[2], s[3], s[4], s[5]])
        for i in range(6):
            grand[i] += s[i]
    if total_rows:
        total_rows.append(["Всего", int(grand[0]), *[round(v, 2) for v in grand[1:]]])

    return {
        "readings": (reading_headers, reading_rows),
        "tariffs": (tariff_headers, tariff_rows),
        "monthly": (month_headers, month_rows),
        "totals": (total_headers, total_rows),
    }


def _rate_for(tariffs: list[Tariff], resource_type: str, period: date) -> float:
    """Resolve the tariff rate effective for the given month."""
    candidates = [
        t
        for t in tariffs
        if t.resource_type == resource_type
        and t.effective_from <= period
        and (t.effective_to is None or t.effective_to >= period)
    ]
    if not candidates:
        return 0.0
    return max(candidates, key=lambda t: t.effective_from).rate


def _month_name(d: date) -> str:
    return MONTHS_RU[d.month - 1]


# ---------------------------------------------------------------- excel

def _format_xlsx_column(ws, column_idx: int, fmt: str, max_row: int) -> None:
    for r in range(2, max_row + 1):
        ws.cell(row=r, column=column_idx).number_format = fmt


def _build_xlsx(tables: dict[str, tuple[list[str], list[list]]]) -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(color="FFFFFF", bold=True)

    specs = [
        ("Показания", "readings"),
        ("Тарифы", "tariffs"),
        ("Расходы по месяцам", "monthly"),
        ("Итоги по годам", "totals"),
    ]

    for sheet_name, key in specs:
        headers, rows = tables[key]
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)
        for row in rows:
            ws.append(row)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")

        ws.freeze_panes = "A2"

        # Числовые форматы: вода/электричество — объём, всё остальное — деньги.
        if key == "readings":
            for col in (4, 5):
                _format_xlsx_column(ws, col, "0.000", ws.max_row)
            _format_xlsx_column(ws, 6, "0.##", ws.max_row)
        elif key == "monthly":
            for col in (4, 5):
                _format_xlsx_column(ws, col, "0.000", ws.max_row)
            _format_xlsx_column(ws, 6, "0.##", ws.max_row)
            for col in range(7, 12):
                _format_xlsx_column(ws, col, "#,##0.00", ws.max_row)
        elif key == "totals":
            for col in range(3, 8):
                _format_xlsx_column(ws, col, "#,##0.00", ws.max_row)
        elif key == "tariffs":
            _format_xlsx_column(ws, 2, "#,##0.00", ws.max_row)

        # Ширина колонок по содержимому.
        for col_cells in ws.columns:
            idx = col_cells[0].column
            length = max(
                (len(str(c.value)) if c.value is not None else 0) for c in col_cells
            )
            ws.column_dimensions[get_column_letter(idx)].width = min(max(length + 3, 10), 60)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- csv

def _csv_bytes(
    headers: list[str],
    rows: list[list],
    delimiter: str = ";",
    money_cols: tuple[int, ...] = (),
) -> bytes:
    """CSV в кодировке utf-8-sig, чтобы Excel корректно открывал кириллицу.

    Разделитель «;» — стандартный для локализованных версий Excel.
    Денежные колонки пишутся строкой с двумя знаками после запятой.
    """

    def cell(value, idx: int):
        if idx in money_cols and isinstance(value, (int, float)):
            return f"{value:.2f}"
        return value

    buf = StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow([cell(v, i) for i, v in enumerate(row)])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


# Колонки с деньгами (0-индекс) для CSV-выгрузок по таблицам.
CSV_MONEY_COLUMNS = {
    "readings": (),
    "tariffs": (1,),
    "monthly": (6, 7, 8, 9, 10),
    "totals": (2, 3, 4, 5, 6),
}


def _csv_file_bytes(tables: dict, key: str) -> bytes:
    headers, rows = tables[key]
    return _csv_bytes(headers, rows, money_cols=CSV_MONEY_COLUMNS[key])


def _csv_response(tables: dict, key: str, filename: str) -> Response:
    content = _csv_file_bytes(tables, key)
    return Response(
        content=content,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------- routes

@router.get(
    "/xlsx",
    summary="Скачать все данные книгой Excel",
    description="Единый файл .xlsx с листами: показания счётчиков, тарифы, "
    "помесячные расходы и итоги по годам.",
)
async def export_xlsx(db: AsyncSession = Depends(get_db)):
    tables = await _collect(db)
    content = _build_xlsx(tables)
    return Response(
        content=content,
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": 'attachment; filename="kvarplata.xlsx"'},
    )


@router.get(
    "/csv/all",
    summary="Скачать все данные в CSV (ZIP-архив)",
    description="Архив .zip, внутри — отдельные CSV-файлы: показания, тарифы, "
    "расходы по месяцам и итоги по годам.",
)
async def export_csv_archive(db: AsyncSession = Depends(get_db)):
    tables = await _collect(db)
    buf = BytesIO()
    with ZipFile(buf, "w", ZIP_DEFLATED) as zf:
        for key, filename in (
            ("readings", "readings.csv"),
            ("tariffs", "tariffs.csv"),
            ("monthly", "monthly.csv"),
            ("totals", "totals.csv"),
        ):
            zf.writestr(filename, _csv_file_bytes(tables, key))
    return Response(
        content=buf.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": 'attachment; filename="kvarplata_csv.zip"'},
    )


@router.get("/csv/readings", summary="Скачать показания в CSV")
async def export_readings_csv(db: AsyncSession = Depends(get_db)):
    return _csv_response(await _collect(db), "readings", "kvarplata_readings.csv")


@router.get("/csv/tariffs", summary="Скачать тарифы в CSV")
async def export_tariffs_csv(db: AsyncSession = Depends(get_db)):
    return _csv_response(await _collect(db), "tariffs", "kvarplata_tariffs.csv")


@router.get("/csv/monthly", summary="Скачать помесячные расходы в CSV")
async def export_monthly_csv(db: AsyncSession = Depends(get_db)):
    return _csv_response(await _collect(db), "monthly", "kvarplata_monthly.csv")


@router.get("/csv/totals", summary="Скачать итоги по годам в CSV")
async def export_totals_csv(db: AsyncSession = Depends(get_db)):
    return _csv_response(await _collect(db), "totals", "kvarplata_totals.csv")
